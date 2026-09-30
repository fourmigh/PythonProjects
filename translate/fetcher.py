import os

import requests

from translator import TranslationError

UA = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
      '(KHTML, like Gecko) Chrome/120.0 Safari/537.36')


def read_file(path):
    if not os.path.exists(path):
        raise TranslationError(f'文件不存在: {path}')
    ext = os.path.splitext(path)[1].lower()
    if ext == '.docx':
        return _read_docx(path)
    if ext == '.pdf':
        return _read_pdf(path)
    return _read_txt(path)


def _read_txt(path):
    for enc in ('utf-8-sig', 'utf-8', 'gbk', 'gb18030'):
        try:
            with open(path, 'r', encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, LookupError):
            continue
    raise TranslationError('无法识别文件编码')


def _read_docx(path):
    try:
        import docx
    except ImportError:
        raise TranslationError('读取 docx 需要安装 python-docx')
    d = docx.Document(path)
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t:
                    lines.append(t)
    return '\n'.join(lines)


def _read_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        raise TranslationError('读取 pdf 需要安装 pypdf')
    reader = PdfReader(path)
    out = []
    for page in reader.pages:
        text = page.extract_text() or ''
        for l in text.split('\n'):
            if l.strip():
                out.append(l.strip())
    return '\n'.join(out)


def fetch_url(url):
    if not url.startswith('http'):
        url = 'http://' + url
    try:
        resp = requests.get(url, headers={'User-Agent': UA}, timeout=30)
        resp.raise_for_status()
    except Exception as e:
        raise TranslationError(f'网页请求失败: {e}')
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        raise TranslationError('网页解析需要安装 beautifulsoup4')
    resp.encoding = resp.apparent_encoding or 'utf-8'
    soup = BeautifulSoup(resp.text, 'html.parser')
    for tag in soup.find_all(['script', 'style', 'noscript', 'header',
                              'footer', 'nav', 'aside', 'iframe', 'form']):
        tag.decompose()
    lines = []
    for para in soup.find_all(['p', 'h1', 'h2', 'h3', 'h4', 'h5',
                               'li', 'td', 'th', 'blockquote', 'pre']):
        text = ' '.join(para.get_text().split())
        if text and text not in lines:
            lines.append(text)
    if not lines:
        body = soup.get_text('\n')
        lines = [l.strip() for l in body.splitlines() if l.strip()]
    return '\n'.join(lines)
